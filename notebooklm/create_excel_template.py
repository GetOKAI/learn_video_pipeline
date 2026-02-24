#!/usr/bin/env python3
"""Generate an Excel template for NotebookLM batch video runs.

Creates a .xlsx file with the required columns and optionally
pre-populates rows from an existing JSON prompts file.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_DIR = SCRIPT_DIR.parent
DEFAULT_OUTPUT = SCRIPT_DIR / "video_queue.xlsx"

# Column definitions — order matters
COLUMNS = [
    ("video_title",        "Video Title",              40),
    ("studio_note",        "Studio Note / Prompt",     80),
    ("notebook_name",      "Notebook Name",            25),
    ("new_notebook",       "New Notebook? (yes/no)",   18),
    ("status",             "Status",                   14),
    ("run_timestamp",      "Run Timestamp",            22),
    ("notebook_url",       "Notebook URL",             50),
    ("error",              "Error (if any)",           40),
]


def create_template(output_path: Path, prompts_json: Path | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Video Queue"

    # ---- Header row ----
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    for col_idx, (_, label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ---- Pre-populate from JSON if provided ----
    if prompts_json and prompts_json.exists():
        items = json.loads(prompts_json.read_text(encoding="utf-8"))
        for row_idx, item in enumerate(items, start=2):
            ws.cell(row=row_idx, column=1, value=item.get("Video Title", ""))
            ws.cell(row=row_idx, column=2, value=item.get("Prompts used for original video", ""))
            ws.cell(row=row_idx, column=3, value="OkAI Video")
            ws.cell(row=row_idx, column=4, value="yes")
            ws.cell(row=row_idx, column=5, value="")       # status — blank = pending
            ws.cell(row=row_idx, column=6, value="")
            ws.cell(row=row_idx, column=7, value="")
            ws.cell(row=row_idx, column=8, value="")

    # Freeze header row
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(str(output_path))
    print(f"✅ Template saved to {output_path}")
    print(f"   Rows (excl. header): {ws.max_row - 1}")


def main() -> int:
    import argparse

    parser = argparse.ArgumentParser(description="Create Excel template for video queue")
    parser.add_argument(
        "--output", "-o",
        default=str(DEFAULT_OUTPUT),
        help=f"Output .xlsx path (default: {DEFAULT_OUTPUT})",
    )
    parser.add_argument(
        "--prompts-json", "-p",
        default=str(REPO_DIR / "Video_generation_prompts_filtered.json"),
        help="Optional JSON file to pre-populate rows from",
    )
    args = parser.parse_args()

    prompts = Path(args.prompts_json) if args.prompts_json else None
    create_template(Path(args.output), prompts)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
