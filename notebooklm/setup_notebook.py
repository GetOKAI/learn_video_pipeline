#!/usr/bin/env python3
"""Setup and launcher for NotebookLM Playwright automation."""

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
AUTOMATION_SCRIPT = SCRIPT_DIR / "notebook_api_generator.py"
DEFAULT_PROFILE_DIR = Path.home() / ".notebooklm_playwright_profile"


def run_cmd(cmd: list[str]) -> int:
    return subprocess.call(cmd)


def install_dependencies() -> int:
    print("Installing Playwright...")
    rc = run_cmd([sys.executable, "-m", "pip", "install", "playwright"])
    if rc != 0:
        return rc

    print("Installing Chromium browser for Playwright...")
    return run_cmd([sys.executable, "-m", "playwright", "install", "chromium"])


def login(profile_dir: str, timeout: int) -> int:
    cmd = [
        sys.executable,
        str(AUTOMATION_SCRIPT),
        "--login-only",
        "--profile-dir",
        profile_dir,
        "--login-timeout",
        str(timeout),
    ]
    return run_cmd(cmd)


def run_workflow(args: argparse.Namespace) -> int:
    cmd = [
        sys.executable,
        str(AUTOMATION_SCRIPT),
        "--video-title",
        args.video_title,
        "--studio-note",
        args.studio_note,
        "--notebook-name",
        args.notebook_name,
        "--notebook-id",
        args.notebook_id,
        "--additional-notes",
        args.additional_notes,
        "--custom-visual-style",
        args.custom_visual_style,
        "--host-focus-notes",
        args.host_focus_notes,
        "--research-query",
        args.research_query,
        "--profile-dir",
        args.profile_dir,
        "--login-timeout",
        str(args.login_timeout),
        "--timeout-ms",
        str(args.timeout_ms),
        "--slow-mo",
        str(args.slow_mo),
        "--keep-open",
        str(args.keep_open),
    ]

    if args.start_url:
        cmd.extend(["--start-url", args.start_url])

    if args.new_notebook:
        cmd.append("--new-notebook")

    if args.output_json:
        cmd.extend(["--output-json", args.output_json])

    if args.debug_dir:
        cmd.extend(["--debug-dir", args.debug_dir])

    if args.disable_api_fallback:
        cmd.append("--disable-api-fallback")

    if args.api_storage_path:
        cmd.extend(["--api-storage-path", args.api_storage_path])

    if args.disable_fast_research:
        cmd.append("--disable-fast-research")

    if args.disable_refresh_before_video_overview:
        cmd.append("--disable-refresh-before-video-overview")

    if args.disable_popup_closing:
        cmd.append("--disable-popup-closing")

    if args.headless:
        cmd.append("--headless")

    return run_cmd(cmd)


def run_batch(config_file: str, profile_dir: str) -> int:
    config_path = Path(config_file)
    if not config_path.exists():
        print(f"Config file not found: {config_path}")
        return 1

    items = json.loads(config_path.read_text(encoding="utf-8"))
    if not isinstance(items, list):
        print("Batch config must be a JSON array")
        return 1

    rc = 0
    for index, item in enumerate(items, start=1):
        title = item.get("video_title") or item.get("title")
        note = item.get("studio_note") or item.get("instructions")
        if not title or not note:
            print(f"Skipping item {index}: missing video_title/title or studio_note/instructions")
            rc = 1
            continue

        cmd = [
            sys.executable,
            str(AUTOMATION_SCRIPT),
            "--video-title",
            title,
            "--studio-note",
            note,
            "--notebook-name",
            item.get("notebook_name", ""),
            "--notebook-id",
            item.get("notebook_id", ""),
            "--additional-notes",
            item.get("additional_notes", ""),
            "--custom-visual-style",
            item.get("custom_visual_style", item.get("additional_notes", "")),
            "--host-focus-notes",
            item.get("host_focus_notes", ""),
            "--research-query",
            item.get("research_query", ""),
            "--profile-dir",
            profile_dir,
            "--login-timeout",
            str(item.get("login_timeout", 180)),
            "--timeout-ms",
            str(item.get("timeout_ms", 25000)),
            "--slow-mo",
            str(item.get("slow_mo", 60)),
            "--keep-open",
            str(item.get("keep_open", 6)),
        ]

        start_url = item.get("start_url")
        if start_url:
            cmd.extend(["--start-url", start_url])

        if item.get("new_notebook"):
            cmd.append("--new-notebook")

        output_json = item.get("output_json")
        if output_json:
            cmd.extend(["--output-json", output_json])

        debug_dir = item.get("debug_dir")
        if debug_dir:
            cmd.extend(["--debug-dir", debug_dir])

        if item.get("disable_api_fallback"):
            cmd.append("--disable-api-fallback")

        api_storage_path = item.get("api_storage_path")
        if api_storage_path:
            cmd.extend(["--api-storage-path", api_storage_path])

        if item.get("disable_fast_research"):
            cmd.append("--disable-fast-research")

        if item.get("disable_refresh_before_video_overview"):
            cmd.append("--disable-refresh-before-video-overview")

        if item.get("disable_popup_closing"):
            cmd.append("--disable-popup-closing")

        if item.get("headless"):
            cmd.append("--headless")

        print(f"\nRunning workflow {index}/{len(items)}: {title}")
        item_rc = run_cmd(cmd)
        if item_rc != 0:
            rc = item_rc

    return rc


# ---------------------------------------------------------------------------
# Excel-based batch runner
# ---------------------------------------------------------------------------

# Expected Excel column headers → internal keys
_EXCEL_COL_MAP = {
    "video title":            "video_title",
    "studio note / prompt":   "studio_note",
    "notebook name":          "notebook_name",
    "new notebook? (yes/no)": "new_notebook",
    "status":                 "status",
    "run timestamp":          "run_timestamp",
    "notebook url":           "notebook_url",
    "error (if any)":         "error",
}

# Statuses that mean "already done — skip"
_DONE_STATUSES = {"done", "success", "completed", "partial", "failed"}

# Persistent history file lives next to setup_notebook.py
_HISTORY_FILE = SCRIPT_DIR / "run_history.json"


# ---------------------------------------------------------------------------
# Persistent run-history helper (the core failsafe)
# ---------------------------------------------------------------------------

class RunHistory:
    """Append-only ledger of every video title + prompt ever dispatched.

    Persisted as a JSON file so it survives across runs, spreadsheets,
    and accidental status-column resets in Excel.
    """

    def __init__(self, path: Path = _HISTORY_FILE) -> None:
        self.path = path
        self._titles: set[str] = set()        # normalised (lowered+stripped) titles
        self._prompt_hashes: set[str] = set()  # SHA-256 of prompt text
        self._entries: list[dict] = []
        self._load()

    # -- persistence ---------------------------------------------------------

    def _load(self) -> None:
        if self.path.exists():
            try:
                data = json.loads(self.path.read_text(encoding="utf-8"))
                self._entries = data.get("entries", [])
                for e in self._entries:
                    self._titles.add(e["title_key"])
                    self._prompt_hashes.add(e["prompt_hash"])
            except Exception:
                pass  # corrupted file — start fresh

    def _save(self) -> None:
        payload = {"entries": self._entries}
        self.path.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    # -- query ---------------------------------------------------------------

    @staticmethod
    def _normalise_title(title: str) -> str:
        return title.strip().lower()

    @staticmethod
    def _hash_prompt(prompt: str) -> str:
        return hashlib.sha256(prompt.strip().encode("utf-8")).hexdigest()

    def was_title_run(self, title: str) -> bool:
        return self._normalise_title(title) in self._titles

    def was_prompt_run(self, prompt: str) -> bool:
        return self._hash_prompt(prompt) in self._prompt_hashes

    def check(self, title: str, prompt: str) -> str | None:
        """Return a human-readable reason if this video was already run, else None."""
        norm_title = self._normalise_title(title)
        p_hash = self._hash_prompt(prompt)
        if norm_title in self._titles:
            return f"title already in history: '{title}'"
        if p_hash in self._prompt_hashes:
            return f"prompt content already in history (hash {p_hash[:12]}…)"
        return None

    # -- record --------------------------------------------------------------

    def record(self, title: str, prompt: str, *, status: str = "done",
               notebook_url: str = "", excel_row: int | None = None) -> None:
        norm_title = self._normalise_title(title)
        p_hash = self._hash_prompt(prompt)
        self._titles.add(norm_title)
        self._prompt_hashes.add(p_hash)
        self._entries.append({
            "title_key": norm_title,
            "title_original": title.strip(),
            "prompt_hash": p_hash,
            "status": status,
            "notebook_url": notebook_url,
            "excel_row": excel_row,
            "timestamp": datetime.now().isoformat(timespec="seconds"),
        })
        self._save()


def _col_index(headers: list[str], key: str) -> int | None:
    """Return 1-based column index for *key* in headers (case-insensitive)."""
    for i, h in enumerate(headers, start=1):
        if h.strip().lower() == key.lower():
            return i
    return None


def run_excel(args: argparse.Namespace) -> int:
    """Read an Excel sheet, run the next *batch_size* pending videos, write status back."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("openpyxl is required. Install with:  pip install openpyxl")
        return 1

    excel_path = Path(args.excel)
    if not excel_path.exists():
        print(f"Excel file not found: {excel_path}")
        return 1

    wb = load_workbook(str(excel_path))
    ws = wb.active

    # --- Parse header row → column mapping --------------------------------
    headers: list[str] = []
    for cell in ws[1]:
        headers.append(str(cell.value or "").strip())

    col = {}
    for header_text, internal_key in _EXCEL_COL_MAP.items():
        idx = _col_index(headers, header_text)
        if idx is not None:
            col[internal_key] = idx

    # Minimum required columns
    if "video_title" not in col or "studio_note" not in col:
        print("Excel sheet must have at least 'Video Title' and 'Studio Note / Prompt' columns.")
        return 1

    status_col = col.get("status")
    ts_col = col.get("run_timestamp")
    url_col = col.get("notebook_url")
    err_col = col.get("error")

    # --- Load persistent history (failsafe layer) -------------------------
    history = RunHistory()
    print(f"📖 Loaded run history: {len(history._entries)} previous run(s) on file.")

    # --- Collect pending rows (skip header row 1) -------------------------
    pending_rows: list[int] = []  # 1-based row numbers
    seen_titles: set[str] = set()

    for row_num in range(2, ws.max_row + 1):
        title_val = ws.cell(row=row_num, column=col["video_title"]).value
        if not title_val or not str(title_val).strip():
            continue  # empty row

        title_str = str(title_val).strip()
        prompt_str = str(ws.cell(row=row_num, column=col["studio_note"]).value or "").strip()

        # --- Layer 1: Excel status column ---------------------------------
        if status_col:
            current_status = str(ws.cell(row=row_num, column=status_col).value or "").strip().lower()
            if current_status in _DONE_STATUSES:
                seen_titles.add(title_str.lower())
                continue

        # --- Layer 2: In-memory title dedup within this scan --------------
        if title_str.lower() in seen_titles:
            print(f"  ⚠ Skipping duplicate title in sheet (row {row_num}): {title_str}")
            if status_col:
                ws.cell(row=row_num, column=status_col, value="skipped-duplicate")
            continue

        # --- Layer 3: Persistent history (title OR prompt content) --------
        dup_reason = history.check(title_str, prompt_str)
        if dup_reason:
            print(f"Skipping row {row_num} — {dup_reason}")
            if status_col:
                ws.cell(row=row_num, column=status_col, value="skipped-history")
            continue

        seen_titles.add(title_str.lower())
        pending_rows.append(row_num)

    if not pending_rows:
        print("✅ No pending videos found in the spreadsheet — all rows processed or empty.")
        wb.save(str(excel_path))
        return 0

    batch_size = args.batch_size
    rows_to_run = pending_rows[:batch_size]

    print(f"\n📋 Found {len(pending_rows)} pending video(s). Running next {len(rows_to_run)}.\n")

    rc = 0
    for seq, row_num in enumerate(rows_to_run, start=1):
        title = str(ws.cell(row=row_num, column=col["video_title"]).value).strip()
        note = str(ws.cell(row=row_num, column=col["studio_note"]).value or "").strip()
        nb_name = str(ws.cell(row=row_num, column=col.get("notebook_name", 0)).value or "").strip() if "notebook_name" in col else ""
        new_nb_val = str(ws.cell(row=row_num, column=col.get("new_notebook", 0)).value or "").strip().lower() if "new_notebook" in col else ""
        want_new_nb = new_nb_val in ("yes", "true", "1", "y")

        if not note:
            print(f"  ⚠ Row {row_num}: empty studio note, skipping.")
            if status_col:
                ws.cell(row=row_num, column=status_col, value="skipped-no-prompt")
            wb.save(str(excel_path))
            continue

        # --- Final pre-flight dedup check (race-condition guard) ----------
        dup_reason = history.check(title, note)
        if dup_reason:
            print(f" Pre-flight block for row {row_num} — {dup_reason}")
            if status_col:
                ws.cell(row=row_num, column=status_col, value="skipped-history")
            wb.save(str(excel_path))
            continue

        # Prepare per-video output JSON for capturing notebook_url
        per_video_json = SCRIPT_DIR / "debug" / f"run_row{row_num}.json"
        per_video_json.parent.mkdir(parents=True, exist_ok=True)

        cmd = [
            sys.executable,
            str(AUTOMATION_SCRIPT),
            "--video-title", title,
            "--studio-note", note,
            "--notebook-name", nb_name or "OkAI Video",
            "--notebook-id", "",
            "--additional-notes", "",
            "--custom-visual-style", "",
            "--host-focus-notes", "",
            "--research-query", "",
            "--profile-dir", args.profile_dir,
            "--login-timeout", str(args.login_timeout),
            "--timeout-ms", str(args.timeout_ms),
            "--slow-mo", str(args.slow_mo),
            "--keep-open", str(args.keep_open),
            "--output-json", str(per_video_json),
        ]

        if want_new_nb:
            cmd.append("--new-notebook")

        if args.debug_dir:
            cmd.extend(["--debug-dir", args.debug_dir])

        if args.headless:
            cmd.append("--headless")

        print(f"\n{'='*60}")
        print(f"  [{seq}/{len(rows_to_run)}]  Row {row_num}: {title}")
        print(f"{'='*60}")

        # Mark as in-progress
        if status_col:
            ws.cell(row=row_num, column=status_col, value="running")
            wb.save(str(excel_path))

        item_rc = run_cmd(cmd)
        now_ts = datetime.now().isoformat(timespec="seconds")

        # Read back the per-video result JSON if it exists
        notebook_url = ""
        error_msg = ""
        final_status = "done" if item_rc == 0 else "failed"
        if per_video_json.exists():
            try:
                result_data = json.loads(per_video_json.read_text(encoding="utf-8"))
                notebook_url = result_data.get("notebook_url", "")
                errors = result_data.get("errors", [])
                if errors:
                    error_msg = "; ".join(errors)
                run_status = result_data.get("status", "")
                if run_status:
                    final_status = run_status
            except Exception:
                pass

        # Write results back to Excel
        if status_col:
            ws.cell(row=row_num, column=status_col, value=final_status)
        if ts_col:
            ws.cell(row=row_num, column=ts_col, value=now_ts)
        if url_col and notebook_url:
            ws.cell(row=row_num, column=url_col, value=notebook_url)
        if err_col and error_msg:
            ws.cell(row=row_num, column=err_col, value=error_msg)

        # Save after every video so progress is never lost
        wb.save(str(excel_path))

        # --- Record in persistent history (the failsafe) ------------------
        history.record(
            title, note,
            status=final_status,
            notebook_url=notebook_url,
            excel_row=row_num,
        )

        print(f"  Row {row_num} → {final_status}")

        if item_rc != 0:
            rc = item_rc

    remaining = len(pending_rows) - len(rows_to_run)
    print(f"\n🏁 Batch complete. {len(rows_to_run)} processed, {remaining} remaining.")
    return rc


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Setup and run NotebookLM Playwright video workflow"
    )
    sub = parser.add_subparsers(dest="command", required=True)

    sub.add_parser("install", help="Install Playwright and Chromium")

    login_parser = sub.add_parser("login", help="Open NotebookLM and complete manual login")
    login_parser.add_argument("--profile-dir", default=str(DEFAULT_PROFILE_DIR))
    login_parser.add_argument("--timeout", type=int, default=240)

    run_parser = sub.add_parser("run", help="Run one video generation workflow")
    run_parser.add_argument("--video-title", required=True)
    run_parser.add_argument("--studio-note", required=True)
    run_parser.add_argument("--notebook-name", default="")
    run_parser.add_argument("--notebook-id", default="")
    run_parser.add_argument("--new-notebook", action="store_true")
    run_parser.add_argument("--additional-notes", default="")
    run_parser.add_argument("--custom-visual-style", default="")
    run_parser.add_argument("--host-focus-notes", default="")
    run_parser.add_argument("--research-query", default="")
    run_parser.add_argument("--start-url", default=None)
    run_parser.add_argument("--profile-dir", default=str(DEFAULT_PROFILE_DIR))
    run_parser.add_argument("--output-json", default="")
    run_parser.add_argument("--login-timeout", type=int, default=180)
    run_parser.add_argument("--timeout-ms", type=int, default=25000)
    run_parser.add_argument("--slow-mo", type=int, default=60)
    run_parser.add_argument("--keep-open", type=int, default=8)
    run_parser.add_argument("--debug-dir", default="")
    run_parser.add_argument("--disable-api-fallback", action="store_true")
    run_parser.add_argument("--api-storage-path", default="")
    run_parser.add_argument("--disable-fast-research", action="store_true")
    run_parser.add_argument("--disable-refresh-before-video-overview", action="store_true")
    run_parser.add_argument("--disable-popup-closing", action="store_true")
    run_parser.add_argument("--headless", action="store_true")

    batch_parser = sub.add_parser("batch", help="Run workflows from JSON config array")
    batch_parser.add_argument("--config", required=True)
    batch_parser.add_argument("--profile-dir", default=str(DEFAULT_PROFILE_DIR))

    excel_parser = sub.add_parser(
        "run-excel",
        help="Run next batch of videos from an Excel spreadsheet (.xlsx)",
    )
    excel_parser.add_argument(
        "--excel", required=True,
        help="Path to the .xlsx file with the video queue",
    )
    excel_parser.add_argument(
        "--batch-size", "-n", type=int, default=10,
        help="Number of pending videos to process (default: 10)",
    )
    excel_parser.add_argument("--profile-dir", default=str(DEFAULT_PROFILE_DIR))
    excel_parser.add_argument("--login-timeout", type=int, default=180)
    excel_parser.add_argument("--timeout-ms", type=int, default=25000)
    excel_parser.add_argument("--slow-mo", type=int, default=60)
    excel_parser.add_argument("--keep-open", type=int, default=8)
    excel_parser.add_argument("--debug-dir", default="")
    excel_parser.add_argument("--headless", action="store_true")

    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    if args.command == "install":
        return install_dependencies()
    if args.command == "login":
        return login(args.profile_dir, args.timeout)
    if args.command == "run":
        return run_workflow(args)
    if args.command == "batch":
        return run_batch(args.config, args.profile_dir)
    if args.command == "run-excel":
        return run_excel(args)

    parser.print_help()
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
